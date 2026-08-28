from pptx import Presentation
prs = Presentation('c:/Users/chan/Documents/semiconductor-ai-project/BG_analysis_output/BG.pptx')
for i, slide in enumerate(prs.slides):
    text = []
    for shape in slide.shapes:
        if hasattr(shape, 'text'):
            text.append(shape.text)
    print(f'Slide {i+1}: ', ' | '.join(text).replace('
', ' '))
